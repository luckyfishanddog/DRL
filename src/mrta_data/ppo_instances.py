"""Fail-closed intake for historical frozen PPO platform workbooks.

Excel support is intentionally imported only inside workbook-facing functions,
so importing :mod:`mrta_reference` or :mod:`mrta_data` does not require it.
No function in this module calls a platform generator or removes a weld row.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
import hashlib
import json
import math
from pathlib import Path
import statistics
from typing import Any, Iterable, Mapping, Sequence


DATASET_MANIFEST_ID = "PPO_DATASET_MANIFEST_V1"
PHASE3_SMOKESET_ID = "PPO_PHASE3_SMOKESET_V1"
PHASE3_SMOKESET_V2_ID = "PPO_PHASE3_SMOKESET_V2"
INIT_BOOTSTRAP_DEVSET_ID = "PPO_INIT_BOOTSTRAP_DEVSET_V1"
CURRENT_SCHEMA_ID = "PPO_FROZEN_FAMILY_SCHEMA_V1"
LEGACY_SCHEMA_ID = "PPO_LEGACY_FROZEN_FAMILY_SCHEMA_V1"
UNRECOGNIZED_SCHEMA_ID = "UNRECOGNIZED"

RAW_COLUMNS = (
    "weld_id",
    "instance_index",
    "group_id",
    "orig_index",
    "x1",
    "y1",
    "z1",
    "x2",
    "y2",
    "z2",
    "length_m",
)
CURRENT_META_COLUMNS = (
    "record_type",
    "key",
    "value",
    "data_role",
    "family_seed",
    "instance_seed",
    "requested_group_count",
    "placed_group_count",
    "full_placement",
    "actual_weld_count",
    "total_original_weld_length",
    "instance_hash",
    "source_excel",
    "source_excel_sha256",
    "generator_version",
    "spacing_m",
    "allow_rotate",
    "min_weld_length_m",
    "weld_z_m",
    "thickness_m",
    "scientific_config_sha256",
    "sheet_name",
    "generation_status",
    "failure_reason",
    "target_N",
    "actual_N",
    "geometry_hash",
    "geometry_content_hash",
    "hard_valid",
    "roundtrip_hash_match",
    "regeneration_hash_match",
)
LEGACY_META_COLUMNS = (
    "record_type",
    "key",
    "value",
    "sheet_name",
    "target_N",
    "actual_N",
    "geometry_hash",
    "geometry_content_hash",
    "requested_group_count",
    "placed_group_count",
    "hard_valid",
    "roundtrip_hash_match",
    "regeneration_hash_match",
)
COORDINATE_COLUMNS = ("x1", "y1", "z1", "x2", "y2", "z2")
WORKSPACE_X = (0.0, 20.0)
WORKSPACE_Y = (0.0, 12.0)
COORDINATE_DECIMALS = 12
Z_TOLERANCE_M = 1.0e-9
XY_TOLERANCE_M = 1.0e-12


class PPOInstanceError(ValueError):
    """Stable fail-closed data error."""

    def __init__(self, code: str, message: str):
        super().__init__(f"{code}: {message}")
        self.code = code


@dataclass(frozen=True)
class SheetInventory:
    name: str
    row_count: int
    columns: tuple[str, ...]


@dataclass(frozen=True)
class WorkbookInspection:
    path: Path
    relative_path: str
    file_size: int
    raw_file_sha256: str
    schema_id: str
    classification: str
    sheets: tuple[SheetInventory, ...]
    family_metadata: Mapping[str, Any]
    sheet_metadata: Mapping[str, Mapping[str, Any]]
    errors: tuple[str, ...] = ()


@dataclass(frozen=True)
class WeldRow:
    weld_id: str
    source_row_ordinal: int
    x1: float
    y1: float
    z1: float
    x2: float
    y2: float
    z2: float

    @property
    def xy_length(self) -> float:
        return math.hypot(self.x2 - self.x1, self.y2 - self.y1)

    @property
    def length_3d(self) -> float:
        return math.dist((self.x1, self.y1, self.z1), (self.x2, self.y2, self.z2))

    @property
    def midpoint(self) -> tuple[float, float, float]:
        return (
            (self.x1 + self.x2) / 2.0,
            (self.y1 + self.y2) / 2.0,
            (self.z1 + self.z2) / 2.0,
        )


@dataclass(frozen=True)
class PPOPlatformInstance:
    instance_id: str
    workbook_path: Path
    relative_path: str
    sheet_name: str
    schema_id: str
    raw_file_sha256: str
    input_units: str
    coordinate_scale_to_m: float
    source_input_units: str | None
    source_row_count: int
    rows: tuple[WeldRow, ...]
    family_metadata: Mapping[str, Any]
    sheet_metadata: Mapping[str, Any]


@dataclass(frozen=True)
class ValidationResult:
    status: str
    errors: tuple[str, ...]
    z_min: float | None
    z_max: float | None


def _openpyxl():
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - environment-dependent
        raise PPOInstanceError(
            "EXCEL_DEPENDENCY_MISSING",
            "PPO Excel intake requires the optional openpyxl dependency",
        ) from exc
    return openpyxl


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _canonical_json(value: Any) -> str:
    return json.dumps(
        value,
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=False,
        allow_nan=False,
    )


def _parse_meta_value(value: Any) -> Any:
    if isinstance(value, str):
        stripped = value.strip()
        if stripped.startswith(("[", "{")):
            try:
                return json.loads(stripped)
            except json.JSONDecodeError:
                return value
    return value


def _relative_path(path: Path, root: Path | None) -> str:
    resolved = path.resolve()
    if root is None:
        return path.name
    try:
        return resolved.relative_to(root.resolve()).as_posix()
    except ValueError as exc:
        raise PPOInstanceError(
            "WORKBOOK_OUTSIDE_PPO_ROOT", f"{resolved} is outside {root.resolve()}"
        ) from exc


def _instance_id(relative_path: str, sheet_name: str) -> str:
    return f"{Path(relative_path).with_suffix('').as_posix()}::{sheet_name}"


def discover_ppo_instances(ppo_root: str | Path) -> tuple[Path, ...]:
    """Return every Excel workbook below ``ppo_root`` in stable order."""

    root = Path(ppo_root).resolve()
    if not root.is_dir():
        raise PPOInstanceError("PPO_ROOT_NOT_FOUND", str(root))
    return tuple(
        sorted(
            (
                path
                for path in root.rglob("*")
                if path.is_file() and path.suffix.lower() in {".xlsx", ".xls"}
            ),
            key=lambda path: path.relative_to(root).as_posix(),
        )
    )


def _schema_from(meta_columns: tuple[str, ...], raw_columns: Sequence[tuple[str, ...]]) -> str:
    if meta_columns == CURRENT_META_COLUMNS and all(columns == RAW_COLUMNS for columns in raw_columns):
        return CURRENT_SCHEMA_ID
    if meta_columns == LEGACY_META_COLUMNS and all(columns == RAW_COLUMNS for columns in raw_columns):
        return LEGACY_SCHEMA_ID
    return UNRECOGNIZED_SCHEMA_ID


def inspect_ppo_workbook(
    path: str | Path, *, ppo_root: str | Path | None = None
) -> WorkbookInspection:
    """Inspect sheets, columns, row counts, metadata and file identity read-only."""

    workbook_path = Path(path).resolve()
    root = None if ppo_root is None else Path(ppo_root).resolve()
    relative = _relative_path(workbook_path, root)
    if workbook_path.suffix.lower() != ".xlsx":
        return WorkbookInspection(
            workbook_path,
            relative,
            workbook_path.stat().st_size,
            _sha256_file(workbook_path),
            UNRECOGNIZED_SCHEMA_ID,
            "UNRECOGNIZED",
            (),
            {},
            {},
            ("UNSUPPORTED_XLS_FORMAT",),
        )
    openpyxl = _openpyxl()
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    inventories: list[SheetInventory] = []
    family: dict[str, Any] = {}
    sheet_metadata: dict[str, dict[str, Any]] = {}
    errors: list[str] = []
    meta_columns: tuple[str, ...] = ()
    try:
        for worksheet in workbook.worksheets:
            row_values = worksheet.iter_rows(values_only=True)
            first = next(row_values, ())
            columns = tuple("" if value is None else str(value) for value in first)
            inventories.append(
                SheetInventory(worksheet.title, sum(1 for _ in row_values), columns)
            )
        if "META" not in workbook.sheetnames:
            errors.append("MISSING_META_SHEET")
        else:
            rows = list(workbook["META"].iter_rows(values_only=True))
            if rows:
                meta_columns = tuple(str(value) for value in rows[0] if value is not None)
                for values in rows[1:]:
                    record = dict(zip(meta_columns, values))
                    if record.get("record_type") == "family":
                        family[str(record.get("key"))] = _parse_meta_value(record.get("value"))
                    elif record.get("record_type") == "sheet" and record.get("sheet_name"):
                        sheet_metadata[str(record["sheet_name"])] = {
                            key: _parse_meta_value(value) for key, value in record.items()
                        }
            else:
                errors.append("EMPTY_META_SHEET")
        raw_columns = [item.columns for item in inventories if item.name != "META"]
        schema_id = _schema_from(meta_columns, raw_columns)
        if schema_id == UNRECOGNIZED_SCHEMA_ID:
            errors.append("UNRECOGNIZED_SCHEMA")
        expected_sheets = {"META", *sheet_metadata}
        if sheet_metadata and set(workbook.sheetnames) != expected_sheets:
            errors.append("SHEET_SET_MISMATCH")
        classification = (
            "FROZEN_PLATFORM_INSTANCE_FAMILY"
            if schema_id != UNRECOGNIZED_SCHEMA_ID
            else "UNRECOGNIZED"
        )
        return WorkbookInspection(
            workbook_path,
            relative,
            workbook_path.stat().st_size,
            _sha256_file(workbook_path),
            schema_id,
            classification,
            tuple(inventories),
            family,
            sheet_metadata,
            tuple(dict.fromkeys(errors)),
        )
    finally:
        workbook.close()


def _float(value: Any, *, field: str, location: str) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise PPOInstanceError("NON_NUMERIC_COORDINATE", f"{location}:{field}") from exc
    if not math.isfinite(result):
        raise PPOInstanceError("NON_FINITE_COORDINATE", f"{location}:{field}")
    return result


def load_ppo_platform_instance(
    path: str | Path,
    sheet_name: str,
    *,
    ppo_root: str | Path | None = None,
    instance_id: str | None = None,
) -> PPOPlatformInstance:
    """Load one frozen platform sheet without generation, filtering or movement."""

    inspection = inspect_ppo_workbook(path, ppo_root=ppo_root)
    if inspection.schema_id == UNRECOGNIZED_SCHEMA_ID:
        raise PPOInstanceError("UNRECOGNIZED_SCHEMA", inspection.relative_path)
    if sheet_name == "META" or sheet_name not in {item.name for item in inspection.sheets}:
        raise PPOInstanceError("INSTANCE_SHEET_NOT_FOUND", sheet_name)

    openpyxl = _openpyxl()
    workbook = openpyxl.load_workbook(inspection.path, read_only=True, data_only=False)
    try:
        return _load_sheet(
            inspection,
            workbook[sheet_name],
            instance_id=instance_id,
        )
    finally:
        workbook.close()


def _load_sheet(
    inspection: WorkbookInspection,
    worksheet,
    *,
    instance_id: str | None = None,
) -> PPOPlatformInstance:
    sheet_name = worksheet.title
    values = worksheet.iter_rows(values_only=False)
    header_cells = next(values, ())
    columns = tuple("" if cell.value is None else str(cell.value) for cell in header_cells)
    if columns != RAW_COLUMNS:
        raise PPOInstanceError("INVALID_WELD_COLUMNS", sheet_name)
    rows: list[WeldRow] = []
    identifier = instance_id or _instance_id(inspection.relative_path, sheet_name)
    for ordinal, cells in enumerate(values):
        if any(cell.data_type == "f" for cell in cells):
            raise PPOInstanceError("FORMULA_FORBIDDEN", f"{sheet_name}:row={ordinal + 2}")
        raw = dict(zip(columns, (cell.value for cell in cells)))
        stable_id = (
            str(raw["weld_id"]).strip()
            if raw.get("weld_id") is not None and str(raw["weld_id"]).strip()
            else f"{identifier}::weld_{ordinal:04d}"
        )
        location = f"{sheet_name}:row={ordinal + 2}"
        coordinates = {
            name: _float(raw.get(name), field=name, location=location)
            for name in COORDINATE_COLUMNS
        }
        rows.append(
            WeldRow(
                stable_id,
                ordinal,
                coordinates["x1"],
                coordinates["y1"],
                coordinates["z1"],
                coordinates["x2"],
                coordinates["y2"],
                coordinates["z2"],
            )
        )
    return PPOPlatformInstance(
        identifier,
        inspection.path,
        inspection.relative_path,
        sheet_name,
        inspection.schema_id,
        inspection.raw_file_sha256,
        "m",
        1.0,
        (
            None
            if inspection.family_metadata.get("input_units") is None
            else str(inspection.family_metadata.get("input_units"))
        ),
        len(rows),
        tuple(rows),
        inspection.family_metadata,
        inspection.sheet_metadata.get(sheet_name, {}),
    )


def validate_ppo_platform_instance(instance: PPOPlatformInstance) -> ValidationResult:
    errors: list[str] = []
    seen_ids: set[str] = set()
    z_values: list[float] = []
    for row in instance.rows:
        if row.weld_id in seen_ids:
            errors.append(f"DUPLICATE_WELD_ID:{row.weld_id}")
        seen_ids.add(row.weld_id)
        coordinates = (row.x1, row.y1, row.z1, row.x2, row.y2, row.z2)
        if not all(math.isfinite(value) for value in coordinates):
            errors.append(f"NON_FINITE_COORDINATE:row={row.source_row_ordinal}")
            continue
        if not (WORKSPACE_X[0] <= row.x1 <= WORKSPACE_X[1]) or not (
            WORKSPACE_X[0] <= row.x2 <= WORKSPACE_X[1]
        ):
            errors.append(f"OUT_OF_BOUNDS_X:row={row.source_row_ordinal}")
        if not (WORKSPACE_Y[0] <= row.y1 <= WORKSPACE_Y[1]) or not (
            WORKSPACE_Y[0] <= row.y2 <= WORKSPACE_Y[1]
        ):
            errors.append(f"OUT_OF_BOUNDS_Y:row={row.source_row_ordinal}")
        dz = abs(row.z2 - row.z1)
        if row.xy_length <= XY_TOLERANCE_M and dz > Z_TOLERANCE_M:
            errors.append(f"UNSUPPORTED_VERTICAL_WELD:row={row.source_row_ordinal}")
        elif dz > Z_TOLERANCE_M:
            errors.append(f"UNSUPPORTED_NONPLANAR_GEOMETRY:row={row.source_row_ordinal}")
        elif row.xy_length <= XY_TOLERANCE_M:
            errors.append(f"ZERO_XY_WELD_LENGTH:row={row.source_row_ordinal}")
        z_values.extend((row.z1, row.z2))
    if z_values and max(z_values) - min(z_values) > Z_TOLERANCE_M:
        errors.append("UNSUPPORTED_NONPLANAR_GEOMETRY:MULTIPLE_Z_PLANES")
    expected = instance.sheet_metadata.get("actual_weld_count")
    if expected is None:
        expected = instance.sheet_metadata.get("actual_N")
    if expected is not None and int(expected) != len(instance.rows):
        errors.append(f"ROW_COUNT_MISMATCH:metadata={int(expected)}:actual={len(instance.rows)}")
    unique_errors = tuple(dict.fromkeys(errors))
    status = "VALID" if not unique_errors else "INVALID"
    return ValidationResult(
        status,
        unique_errors,
        min(z_values) if z_values else None,
        max(z_values) if z_values else None,
    )


def to_parent_welds(instance: PPOPlatformInstance):
    """Project validated planar rows to the current formal XY model."""

    validation = validate_ppo_platform_instance(instance)
    if validation.status != "VALID":
        raise PPOInstanceError("INVALID_PLATFORM_INSTANCE", ";".join(validation.errors))
    from mrta_reference.model import ParentWeld

    return tuple(
        ParentWeld(row.weld_id, (row.x1, row.y1), (row.x2, row.y2))
        for row in instance.rows
    )


def _canonical_segment(row: WeldRow) -> tuple[float, ...]:
    start = tuple(round(value, COORDINATE_DECIMALS) for value in (row.x1, row.y1, row.z1))
    end = tuple(round(value, COORDINATE_DECIMALS) for value in (row.x2, row.y2, row.z2))
    return (*min(start, end), *max(start, end))


def instance_geometry_hash(instance: PPOPlatformInstance) -> str:
    payload = sorted(_canonical_segment(row) for row in instance.rows)
    return hashlib.sha256(_canonical_json(payload).encode("utf-8")).hexdigest()


def _crosses(a: float, b: float, boundary: float, epsilon: float = 1.0e-9) -> bool:
    return (a - boundary) * (b - boundary) < -(epsilon * epsilon)


def geometry_metrics(instance: PPOPlatformInstance) -> dict[str, float | int]:
    rows = instance.rows
    xs = [value for row in rows for value in (row.x1, row.x2)]
    ys = [value for row in rows for value in (row.y1, row.y2)]
    x_coverage = (max(xs) - min(xs)) / 20.0 if xs else 0.0
    y_coverage = (max(ys) - min(ys)) / 12.0 if ys else 0.0
    midpoints = [row.midpoint for row in rows]
    if midpoints:
        mean_x = statistics.fmean(point[0] for point in midpoints)
        mean_y = statistics.fmean(point[1] for point in midpoints)
        dispersion = math.sqrt(
            statistics.fmean(
                (point[0] - mean_x) ** 2 + (point[1] - mean_y) ** 2
                for point in midpoints
            )
        ) / math.hypot(20.0, 12.0)
    else:
        dispersion = 0.0
    occupied = {
        (
            min(3, max(0, int(point[0] / 20.0 * 4))),
            min(3, max(0, int(point[1] / 12.0 * 4))),
        )
        for point in midpoints
    }
    quadrant_lengths = [0.0, 0.0, 0.0, 0.0]
    lengths: list[float] = []
    cross_y = 0
    cross_x = 0
    for row in rows:
        length = row.xy_length
        lengths.append(length)
        midpoint = row.midpoint
        quadrant = (0 if midpoint[1] >= 6.0 else 2) + (0 if midpoint[0] < 10.0 else 1)
        quadrant_lengths[quadrant] += length
        cross_y += int(_crosses(row.y1, row.y2, 6.0))
        cross_x += int(_crosses(row.x1, row.x2, 10.0))
    quadrant_mean = statistics.fmean(quadrant_lengths)
    quadrant_std = statistics.pstdev(quadrant_lengths)
    total = sum(lengths)
    upper = quadrant_lengths[0] + quadrant_lengths[1]
    lower = quadrant_lengths[2] + quadrant_lengths[3]
    left = quadrant_lengths[0] + quadrant_lengths[2]
    right = quadrant_lengths[1] + quadrant_lengths[3]
    return {
        "bbox_area_ratio": x_coverage * y_coverage,
        "x_coverage_ratio": x_coverage,
        "y_coverage_ratio": y_coverage,
        "midpoint_dispersion": dispersion,
        "grid_occupancy_ratio": len(occupied) / 16.0,
        "quadrant_length_cv": quadrant_std / quadrant_mean if quadrant_mean else 0.0,
        "upper_lower_length_imbalance": abs(upper - lower) / total if total else 0.0,
        "left_right_length_imbalance": abs(left - right) / total if total else 0.0,
        "cross_y6_count": cross_y,
        "cross_x10_count": cross_x,
        "total_weld_length_m": total,
        "weld_length_mean": statistics.fmean(lengths) if lengths else 0.0,
        "weld_length_std": statistics.pstdev(lengths) if len(lengths) > 1 else 0.0,
        "weld_length_min": min(lengths, default=0.0),
        "weld_length_max": max(lengths, default=0.0),
    }


def formal_compatibility_metrics(instance: PPOPlatformInstance) -> dict[str, int | str]:
    from mrta_reference.geometry import (
        blocks_for_pattern,
        generate_y_split_patterns,
        robot_is_eligible,
        whole_eligible_rails,
    )
    from mrta_reference.model import Rail, ScientificConfig

    config = ScientificConfig()
    counts: Counter[str] = Counter()
    for parent in to_parent_welds(instance):
        rails = whole_eligible_rails(parent.start, parent.end, config)
        if rails == frozenset({Rail.UPPER}):
            counts["n_whole_upper_only"] += 1
        elif rails == frozenset({Rail.LOWER}):
            counts["n_whole_lower_only"] += 1
        elif rails == frozenset({Rail.UPPER, Rail.LOWER}):
            counts["n_whole_both_rails"] += 1
        else:
            legal = [
                pattern
                for pattern in generate_y_split_patterns(parent, config)
                if all(
                    any(robot_is_eligible(block, robot, config) for robot in range(4))
                    for block in blocks_for_pattern(parent, pattern, config)
                )
            ]
            if legal:
                counts["n_mandatory_y_split"] += 1
            else:
                counts["n_no_legal_pattern"] += 1
        if _crosses(parent.start[1], parent.end[1], 6.0):
            counts["n_cross_By"] += 1
    result: dict[str, int | str] = {
        key: counts[key]
        for key in (
            "n_whole_upper_only",
            "n_whole_lower_only",
            "n_whole_both_rails",
            "n_mandatory_y_split",
            "n_no_legal_pattern",
            "n_cross_By",
        )
    }
    result["formal_compatibility"] = (
        "PASS" if counts["n_no_legal_pattern"] == 0 else "FAIL"
    )
    return result


def _manifest_entry(instance: PPOPlatformInstance) -> dict[str, Any]:
    validation = validate_ppo_platform_instance(instance)
    metrics = geometry_metrics(instance)
    formal: dict[str, Any] = {
        "n_whole_upper_only": 0,
        "n_whole_lower_only": 0,
        "n_whole_both_rails": 0,
        "n_mandatory_y_split": 0,
        "n_no_legal_pattern": 0,
        "n_cross_By": 0,
        "formal_compatibility": "NOT_EVALUATED",
    }
    errors = list(validation.errors)
    if validation.status == "VALID":
        formal = formal_compatibility_metrics(instance)
        if formal["formal_compatibility"] != "PASS":
            errors.append("FORMAL_NO_LEGAL_PATTERN")
    status = "VALID" if not errors else "INVALID"
    return {
        "instance_id": instance.instance_id,
        "relative_path": instance.relative_path,
        "sheet_name": instance.sheet_name,
        "raw_file_sha256": instance.raw_file_sha256,
        "schema_id": instance.schema_id,
        "input_units": instance.input_units,
        "source_input_units": instance.source_input_units,
        "coordinate_scale_to_m": instance.coordinate_scale_to_m,
        "source_row_count": instance.source_row_count,
        "actual_weld_count": len(instance.rows),
        "instance_geometry_hash": instance_geometry_hash(instance),
        "z_min": validation.z_min,
        "z_max": validation.z_max,
        "validation_status": status,
        "validation_errors": errors,
        "duplicate_of": None,
        "duplicate_reasons": [],
        **metrics,
        **formal,
    }


def _distribution(values: Sequence[float | int]) -> dict[str, float | int | None]:
    if not values:
        return {"min": None, "median": None, "max": None}
    return {
        "min": min(values),
        "median": statistics.median(values),
        "max": max(values),
    }


def build_dataset_manifest(ppo_root: str | Path) -> dict[str, Any]:
    """Inspect and validate every recognized frozen platform sheet."""

    root = Path(ppo_root).resolve()
    paths = discover_ppo_instances(root)
    inspections = [inspect_ppo_workbook(path, ppo_root=root) for path in paths]
    inventory: list[dict[str, Any]] = []
    entries: list[dict[str, Any]] = []
    for inspection in inspections:
        inventory.append(
            {
                "relative_path": inspection.relative_path,
                "filename": inspection.path.name,
                "file_size": inspection.file_size,
                "raw_file_sha256": inspection.raw_file_sha256,
                "schema_id": inspection.schema_id,
                "classification": inspection.classification,
                "sheets": [
                    {
                        "sheet_name": sheet.name,
                        "row_count": sheet.row_count,
                        "column_names": list(sheet.columns),
                    }
                    for sheet in inspection.sheets
                ],
                "inspection_errors": list(inspection.errors),
            }
        )
        if inspection.schema_id == UNRECOGNIZED_SCHEMA_ID:
            continue
        openpyxl = _openpyxl()
        workbook = openpyxl.load_workbook(
            inspection.path, read_only=True, data_only=False
        )
        try:
            for sheet in inspection.sheets:
                if sheet.name == "META":
                    continue
                try:
                    instance = _load_sheet(inspection, workbook[sheet.name])
                    entries.append(_manifest_entry(instance))
                except PPOInstanceError as exc:
                    identifier = _instance_id(inspection.relative_path, sheet.name)
                    entries.append(
                        {
                            "instance_id": identifier,
                            "relative_path": inspection.relative_path,
                            "sheet_name": sheet.name,
                            "raw_file_sha256": inspection.raw_file_sha256,
                            "schema_id": inspection.schema_id,
                            "input_units": "m",
                            "source_input_units": inspection.family_metadata.get("input_units"),
                            "coordinate_scale_to_m": 1.0,
                            "source_row_count": sheet.row_count,
                            "actual_weld_count": sheet.row_count,
                            "instance_geometry_hash": None,
                            "total_weld_length_m": None,
                            "z_min": None,
                            "z_max": None,
                            "validation_status": "INVALID",
                            "validation_errors": [exc.code],
                            "duplicate_of": None,
                            "duplicate_reasons": [],
                        }
                    )
        finally:
            workbook.close()

    entries.sort(key=lambda item: item["instance_id"])
    raw_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    geometry_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for entry in entries:
        raw_groups[str(entry["raw_file_sha256"])].append(entry)
        if entry["validation_status"] == "VALID" and entry.get("instance_geometry_hash"):
            geometry_groups[str(entry["instance_geometry_hash"])].append(entry)
    duplicate_raw_files = {
        digest: sorted({entry["relative_path"] for entry in group})
        for digest, group in raw_groups.items()
        if len({entry["relative_path"] for entry in group}) > 1
    }
    duplicate_geometry: dict[str, list[str]] = {}
    for digest, group in geometry_groups.items():
        ordered = sorted(group, key=lambda item: item["instance_id"])
        if len(ordered) <= 1:
            continue
        canonical = ordered[0]["instance_id"]
        duplicate_geometry[digest] = [item["instance_id"] for item in ordered]
        for item in ordered[1:]:
            item["duplicate_of"] = canonical
            item["duplicate_reasons"].append("IDENTICAL_NORMALIZED_GEOMETRY")
    for digest, paths_for_digest in duplicate_raw_files.items():
        canonical_path = min(paths_for_digest)
        for entry in raw_groups[digest]:
            if entry["relative_path"] != canonical_path:
                entry["duplicate_reasons"].append("IDENTICAL_RAW_FILE_SHA256")

    usable = [
        entry
        for entry in entries
        if entry["validation_status"] == "VALID" and entry["duplicate_of"] is None
    ]
    hash_payload = sorted(usable, key=lambda item: item["instance_id"])
    manifest_hash = hashlib.sha256(
        _canonical_json(hash_payload).encode("utf-8")
    ).hexdigest()
    invalid_reasons = Counter(
        error.split(":", 1)[0]
        for entry in entries
        if entry["validation_status"] != "VALID"
        for error in entry["validation_errors"]
    )
    schema_counts = Counter(item["schema_id"] for item in inventory)
    classifications = Counter(item["classification"] for item in inventory)
    metric_names = (
        "bbox_area_ratio",
        "x_coverage_ratio",
        "y_coverage_ratio",
        "midpoint_dispersion",
        "grid_occupancy_ratio",
        "quadrant_length_cv",
        "upper_lower_length_imbalance",
        "left_right_length_imbalance",
        "weld_length_mean",
        "weld_length_std",
        "weld_length_min",
        "weld_length_max",
    )
    return {
        "dataset_manifest_id": DATASET_MANIFEST_ID,
        "dataset_manifest_hash": manifest_hash,
        "identity_policy": (
            "SHA256(canonical JSON of VALID unique entries sorted by instance_id)"
        ),
        "ppo_root_policy": "runtime-only; manifest paths are PPO-root-relative",
        "coordinate_policy": (
            "Frozen V9 raw sheets use metre-valued x/y/z and length_m columns; "
            "META input_units describes the pre-packing source workbook."
        ),
        "projection_policy": (
            "Only single-plane z1~=z2 welds are projected to XY; no row is filtered, "
            "clipped, moved, or regenerated."
        ),
        "formal_compatibility_policy": (
            "ScientificConfig(); n_cross_By counts strict crossings of y=6.0 m"
        ),
        "inventory_summary": {
            "excel_file_count": len(inventory),
            "schema_counts": dict(sorted(schema_counts.items())),
            "classification_counts": dict(sorted(classifications.items())),
            "raw_source_assembly_count": classifications.get("RAW_SOURCE_ASSEMBLY", 0),
            "frozen_platform_family_count": classifications.get(
                "FROZEN_PLATFORM_INSTANCE_FAMILY", 0
            ),
            "unrecognized_file_count": classifications.get("UNRECOGNIZED", 0),
            "platform_instance_sheet_count": len(entries),
        },
        "validation_summary": {
            "valid_instance_count": sum(
                entry["validation_status"] == "VALID" for entry in entries
            ),
            "invalid_instance_count": sum(
                entry["validation_status"] != "VALID" for entry in entries
            ),
            "valid_unique_instance_count": len(usable),
            "invalid_reason_distribution": dict(sorted(invalid_reasons.items())),
            "weld_count_distribution": _distribution(
                [entry["actual_weld_count"] for entry in usable]
            ),
            "geometry_metric_distributions": {
                name: _distribution([entry[name] for entry in usable])
                for name in metric_names
            },
            "formal_compatibility_totals": {
                name: sum(int(entry.get(name, 0)) for entry in entries)
                for name in (
                    "n_whole_upper_only",
                    "n_whole_lower_only",
                    "n_whole_both_rails",
                    "n_mandatory_y_split",
                    "n_no_legal_pattern",
                    "n_cross_By",
                )
            },
        },
        "duplicate_analysis": {
            "identical_raw_file_sha256": duplicate_raw_files,
            "identical_normalized_geometry": duplicate_geometry,
        },
        "workbook_inventory": inventory,
        "instances": entries,
    }


def select_phase3_smokeset(manifest: Mapping[str, Any]) -> dict[str, Any]:
    valid = [
        dict(entry)
        for entry in manifest["instances"]
        if entry["validation_status"] == "VALID" and entry.get("duplicate_of") is None
    ]
    if len(valid) < 3:
        raise PPOInstanceError(
            "INSUFFICIENT_VALID_UNIQUE_INSTANCES", f"found {len(valid)}, require at least 3"
        )
    ordered = sorted(
        valid, key=lambda entry: (entry["actual_weld_count"], entry["instance_geometry_hash"])
    )
    small = ordered[0]
    remaining = [entry for entry in valid if entry["instance_id"] != small["instance_id"]]
    large = min(
        remaining,
        key=lambda entry: (-entry["actual_weld_count"], entry["instance_geometry_hash"]),
    )
    remaining = [entry for entry in remaining if entry["instance_id"] != large["instance_id"]]
    median_n = statistics.median(entry["actual_weld_count"] for entry in valid)
    medium = min(
        remaining,
        key=lambda entry: (
            abs(entry["actual_weld_count"] - median_n),
            entry["instance_geometry_hash"],
        ),
    )

    def record(tier: str, entry: Mapping[str, Any]) -> dict[str, Any]:
        return {
            "tier": tier,
            "instance_id": entry["instance_id"],
            "relative_path": entry["relative_path"],
            "sheet_name": entry["sheet_name"],
            "actual_weld_count": entry["actual_weld_count"],
            "instance_geometry_hash": entry["instance_geometry_hash"],
        }

    return {
        "smokeset_id": PHASE3_SMOKESET_ID,
        "dataset_manifest_id": manifest["dataset_manifest_id"],
        "dataset_manifest_hash": manifest["dataset_manifest_hash"],
        "selection_policy": (
            "small=min N; medium=closest to median N; large=max N; ties by "
            "instance_geometry_hash; selections distinct"
        ),
        "median_weld_count": median_n,
        "instances": [record("small", small), record("medium", medium), record("large", large)],
    }


def _valid_unique_entries(manifest: Mapping[str, Any]) -> list[dict[str, Any]]:
    return [
        dict(entry)
        for entry in manifest["instances"]
        if entry["validation_status"] == "VALID" and entry.get("duplicate_of") is None
    ]


def main_experimental_range_summary(
    manifest: Mapping[str, Any], *, maximum_weld_count: int = 90
) -> dict[str, Any]:
    valid = _valid_unique_entries(manifest)
    in_range = [entry for entry in valid if entry["actual_weld_count"] <= maximum_weld_count]
    out_of_range = [entry for entry in valid if entry["actual_weld_count"] > maximum_weld_count]
    counts = Counter(int(entry["actual_weld_count"]) for entry in in_range)
    return {
        "maximum_weld_count": maximum_weld_count,
        "valid_unique_total": len(valid),
        "valid_unique_in_main_range": len(in_range),
        "valid_unique_out_of_main_range": len(out_of_range),
        "out_of_main_range_percentage": (
            100.0 * len(out_of_range) / len(valid) if valid else 0.0
        ),
        "weld_count_distribution": {
            str(count): counts.get(count, 0)
            for count in range(10, maximum_weld_count + 1)
        },
        "out_of_range_policy": (
            "retained as VALID manifest entries and stress instances; excluded from "
            "the Phase 3 main-range gate"
        ),
    }


def select_phase3_smokeset_v2(manifest: Mapping[str, Any]) -> dict[str, Any]:
    valid = [
        entry
        for entry in _valid_unique_entries(manifest)
        if entry["actual_weld_count"] <= 90
    ]
    if len(valid) < 3:
        raise PPOInstanceError(
            "INSUFFICIENT_MAIN_RANGE_INSTANCES", f"found {len(valid)}, require at least 3"
        )
    ordered = sorted(
        valid,
        key=lambda entry: (entry["actual_weld_count"], entry["instance_geometry_hash"]),
    )
    small = ordered[0]
    remaining = [entry for entry in valid if entry["instance_id"] != small["instance_id"]]
    large = min(
        remaining,
        key=lambda entry: (-entry["actual_weld_count"], entry["instance_geometry_hash"]),
    )
    remaining = [entry for entry in remaining if entry["instance_id"] != large["instance_id"]]
    median_n = statistics.median(entry["actual_weld_count"] for entry in valid)
    medium = min(
        remaining,
        key=lambda entry: (
            abs(entry["actual_weld_count"] - median_n),
            entry["instance_geometry_hash"],
        ),
    )

    def record(tier: str, entry: Mapping[str, Any]) -> dict[str, Any]:
        return {
            "tier": tier,
            "instance_id": entry["instance_id"],
            "relative_path": entry["relative_path"],
            "sheet_name": entry["sheet_name"],
            "actual_weld_count": entry["actual_weld_count"],
            "instance_geometry_hash": entry["instance_geometry_hash"],
        }

    return {
        "smokeset_id": PHASE3_SMOKESET_V2_ID,
        "dataset_manifest_id": manifest["dataset_manifest_id"],
        "dataset_manifest_hash": manifest["dataset_manifest_hash"],
        "main_range_maximum_weld_count": 90,
        "selection_policy": (
            "among VALID unique N<=90: small=min N; medium=closest to median N; "
            "large=max N; ties by instance_geometry_hash; selections distinct"
        ),
        "median_weld_count": median_n,
        "instances": [
            record("small", small),
            record("medium", medium),
            record("large", large),
        ],
    }


def select_init_bootstrap_devset(
    manifest: Mapping[str, Any],
    historical_smokeset: Mapping[str, Any],
    *,
    additional_consumed_smokesets: Sequence[Mapping[str, Any]] = (),
    per_stratum: int = 5,
) -> dict[str, Any]:
    if per_stratum < 1:
        raise ValueError("per_stratum must be positive")
    strata = (
        ("N10_30", 10, 30),
        ("N31_50", 31, 50),
        ("N51_70", 51, 70),
        ("N71_90", 71, 90),
    )
    valid = _valid_unique_entries(manifest)
    selected: list[dict[str, Any]] = []
    stratum_records: list[dict[str, Any]] = []
    globally_used_workbooks: set[str] = set()
    for stratum_id, lower, upper in strata:
        candidates = sorted(
            (
                entry
                for entry in valid
                if lower <= entry["actual_weld_count"] <= upper
            ),
            key=lambda entry: (entry["instance_geometry_hash"], entry["instance_id"]),
        )
        chosen: list[dict[str, Any]] = []
        for entry in candidates:
            if entry["relative_path"] in globally_used_workbooks:
                continue
            chosen.append(entry)
            globally_used_workbooks.add(entry["relative_path"])
            if len(chosen) == per_stratum:
                break
        if len(chosen) < per_stratum:
            chosen_ids = {entry["instance_id"] for entry in chosen}
            for entry in candidates:
                if entry["instance_id"] in chosen_ids:
                    continue
                chosen.append(entry)
                if len(chosen) == per_stratum:
                    break
        if len(chosen) < per_stratum:
            raise PPOInstanceError(
                "INSUFFICIENT_BOOTSTRAP_STRATUM",
                f"{stratum_id}: found {len(chosen)}, require {per_stratum}",
            )
        records = [
            {
                "stratum": stratum_id,
                "instance_id": entry["instance_id"],
                "relative_path": entry["relative_path"],
                "sheet_name": entry["sheet_name"],
                "actual_weld_count": entry["actual_weld_count"],
                "instance_geometry_hash": entry["instance_geometry_hash"],
            }
            for entry in chosen
        ]
        selected.extend(records)
        stratum_records.append(
            {
                "stratum": stratum_id,
                "minimum_weld_count": lower,
                "maximum_weld_count": upper,
                "instances": records,
            }
        )
    smoke_workbooks = {
        str(entry["relative_path"])
        for entry in historical_smokeset["instances"]
    }
    for smokeset in additional_consumed_smokesets:
        smoke_workbooks.update(
            str(entry["relative_path"]) for entry in smokeset["instances"]
        )
    consumed = sorted(
        {
            entry["relative_path"] for entry in selected
        }
        | smoke_workbooks
    )
    return {
        "devset_id": INIT_BOOTSTRAP_DEVSET_ID,
        "dataset_manifest_id": manifest["dataset_manifest_id"],
        "dataset_manifest_hash": manifest["dataset_manifest_hash"],
        "selection_policy": (
            "solver-independent four strata N=10..30,31..50,51..70,71..90; "
            "five per stratum; stable instance_geometry_hash order; prefer distinct "
            "workbooks and use at most one per workbook when possible"
        ),
        "future_split_isolation_policy": (
            "all instances from a development-consumed workbook/generation-seed family "
            "are excluded from untouched TEST"
        ),
        "per_stratum": per_stratum,
        "instances": selected,
        "strata": stratum_records,
        "development_consumed_workbooks": consumed,
        "historical_smoke_workbooks": sorted(
            {str(entry["relative_path"]) for entry in historical_smokeset["instances"]}
        ),
        "all_consumed_smoke_workbooks": sorted(smoke_workbooks),
    }


def write_json(path: str | Path, payload: Mapping[str, Any]) -> None:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False)
        + "\n",
        encoding="utf-8",
    )
